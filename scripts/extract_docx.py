# 📄 本知识库由 QFUD（驻青四校联合辩论培训计划）整理发布，
# 采用 知识共享署名-非商业性使用-相同方式共享 4.0 国际 (CC BY-NC-SA 4.0) 许可。
# 详情：https://creativecommons.org/licenses/by-nc-sa/4.0/deed.zh

import os, sys, re
from pathlib import Path

raw_dir = Path(r"./materials/raw")
out_dir = Path(r"./materials/extracted")
out_dir.mkdir(parents=True, exist_ok=True)

from docx import Document

for docx_file in sorted(raw_dir.rglob("*.docx")):
    out_path = out_dir / (docx_file.stem + ".md")
    if out_path.exists():
        print(f"SKIP: {docx_file.name}")
        continue
    try:
        doc = Document(str(docx_file))
        lines = []
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                lines.append("")
                continue
            style = p.style.name if p.style else ""
            # Heading detection
            if style and "Heading" in style or "heading" in style:
                level = re.search(r'(\d+)', style)
                lvl = int(level.group(1)) if level else 2
                lines.append("#" * min(lvl, 4) + " " + text)
            elif style and ("Title" in style or "title" in style):
                lines.append("# " + text)
            else:
                lines.append(text)
        
        # Also extract tables
        for t_idx, table in enumerate(doc.tables):
            lines.append(f"\n--- 表格 {t_idx+1} ---")
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                lines.append(" | ".join(cells))
            lines.append("")
        
        content = "\n".join(lines)
        out_path.write_text(content, encoding="utf-8")
        print(f"OK  {docx_file.name}  ->  {len(content)} 字符")
    except Exception as e:
        print(f"ERR {docx_file.name}: {e}")

print("\nDONE")